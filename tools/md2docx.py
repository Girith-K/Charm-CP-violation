# turns a markdown report into a Word document, needs python-docx

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

if len(sys.argv) != 3:
    raise SystemExit("usage: md2docx.py REPORT.md OUT.docx")
SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])

doc = Document()

for sec in doc.sections:
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin = sec.right_margin = Inches(1.0)
    sec.top_margin = sec.bottom_margin = Inches(1.0)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(10.5)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, before, after, color in [
    ("Heading 1", 14, 14, 6, "000000"),
    ("Heading 2", 12, 12, 4, "000000"),
    ("Heading 3", 11, 10, 3, "333333"),
]:
    h = doc.styles[name]
    h.font.name = "Times New Roman"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor.from_string(color)
    h.paragraph_format.space_before = Pt(before)
    h.paragraph_format.space_after = Pt(after)


def add_runs(par, text):
    token = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+`)")
    for piece in token.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2])
            r.bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            r.italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        else:
            par.add_run(piece)


def add_table(rows):
    header, body = rows[0], rows[2:]
    t = doc.add_table(rows=len(body) + 1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    tbl = t._tbl
    for borders in tbl.iter(qn("w:tcBorders")):
        borders.getparent().remove(borders)
    for i, row in enumerate([header] + body):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, cell.strip())
            if i == 0:
                for r in p.runs:
                    r.bold = True
            for r in p.runs:
                r.font.size = Pt(9.5)
    return t


def formula(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(text.strip())
    r.font.name = "Cambria Math"
    r.font.size = Pt(10.5)


lines = SRC.read_text(encoding="utf-8").splitlines()


def is_table_start(idx):
    return (idx + 1 < len(lines)
            and lines[idx].startswith("|")
            and re.match(r"^\|[\s:|-]+\|?\s*$", lines[idx + 1]))


i = 0
first_h1_done = False
while i < len(lines):
    ln = lines[i]

    if ln.startswith("```"):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].startswith("```"):
            buf.append(lines[i])
            i += 1
        for b in buf:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(b.strip())
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        i += 1
        continue

    if ln.strip() == "---" or not ln.strip():
        i += 1
        continue

    m_img = re.match(r"^!\[(.*)\]\((.*)\)\s*$", ln)
    if m_img:
        caption, rel = m_img.group(1), m_img.group(2)
        img = SRC.parent / rel
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        run = p.add_run()
        run.add_picture(str(img), width=Inches(6.2))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cap.paragraph_format.space_after = Pt(10)
        mfig = re.match(r"^(Figure \d+:)\s*(.*)$", caption)
        if mfig:
            r = cap.add_run(mfig.group(1) + " ")
            r.bold = True
            add_runs(cap, mfig.group(2))
        else:
            add_runs(cap, caption)
        for r in cap.runs:
            r.font.size = Pt(9)
        i += 1
        continue

    if is_table_start(i):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append([c for c in lines[i].strip().strip("|").split("|")])
            i += 1
        add_table(rows)
        doc.add_paragraph()
        continue

    if ln.startswith("    ") and ln.strip():
        buf = []
        while i < len(lines) and lines[i].startswith("    ") and lines[i].strip():
            buf.append(lines[i].strip())
            i += 1
        for b in buf:
            formula(b)
        continue

    m = re.match(r"^(#{1,3}) (.*)$", ln)
    if m:
        level = len(m.group(1))
        text = m.group(2)
        if level == 1 and not first_h1_done:
            first_h1_done = True
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(16)
        else:
            doc.add_heading(text, level=level)
        i += 1
        continue

    if re.match(r"^\d+\. |^- ", ln):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        add_runs(p, ln)
        i += 1
        continue

    buf = [ln]
    i += 1
    is_ref = re.match(r"^\[\d+\] ", ln)
    while (i < len(lines) and lines[i].strip()
           and not re.match(r"^(#{1,3} |- |\d+\. |    |```|!\[|---$)", lines[i])
           and not is_table_start(i)
           and not re.match(r"^\[\d+\] ", lines[i])):
        buf.append(lines[i])
        i += 1
    p = doc.add_paragraph()
    add_runs(p, " ".join(x.strip() for x in buf))

doc.save(OUT)
print(f"wrote {OUT}")

d2 = Document(OUT)
print(f"paragraphs: {len(d2.paragraphs)}, tables: {len(d2.tables)}")
